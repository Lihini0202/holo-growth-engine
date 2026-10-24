"""Signup export ingestion (CSV or Excel) and column mapping onto RawLead.

Headers are normalised and matched against a synonym table. Unmapped columns are
reported rather than silently dropped.
"""

from __future__ import annotations

import csv
import hashlib
import io
import re
from dataclasses import dataclass, field
from pathlib import Path

from holo_growth.models import RawLead

SYNONYMS: dict[str, tuple[str, ...]] = {
    "full_name": ("name", "full name", "fullname", "your name", "attendee name", "participant"),
    "first_name": ("first name", "firstname", "given name", "fname"),
    "last_name": ("last name", "lastname", "surname", "family name", "lname"),
    "email": ("email", "e mail", "email address", "email id", "mail", "your email"),
    "phone": ("phone", "mobile", "phone number", "contact number", "whatsapp", "tel", "mobile no"),
    "organisation": (
        "university",
        "uni",
        "campus",
        "company",
        "workplace",
        "organisation",
        "organization",
        "institute",
        "school",
    ),
    "city": ("city", "town", "location", "home town", "hometown", "district"),
    "age": ("age", "your age", "age years"),
    "heard_from": ("how did you hear", "how did you hear about us", "source", "referral", "heard from"),
    "reason": (
        "what brings you",
        "what brings you to holo",
        "why",
        "reason",
        "comments",
        "message",
        "anything else",
        "interest",
    ),
    "marketing_opt_in": (
        "send me updates",
        "send me holo updates",
        "marketing consent",
        "opt in",
        "newsletter",
        "consent",
        "keep me updated",
        "subscribe",
    ),
}


def _norm_header(h: str) -> str:
    h = h.strip().lower()
    h = re.sub(r"[^a-z0-9]+", " ", h)
    return re.sub(r"\s+", " ", h).strip()


def map_headers(headers: list[str]) -> tuple[dict[str, str], list[str]]:
    """Return ({original_header: field}, [unmapped headers])."""
    mapping: dict[str, str] = {}
    unmapped: list[str] = []
    taken: set[str] = set()
    for original in headers:
        norm = _norm_header(original)
        match = None
        for fld, syns in SYNONYMS.items():
            if fld in taken:
                continue
            if norm == fld.replace("_", " ") or norm in syns:
                match = fld
                break
        if match is None:  # Second pass: prefix match (e.g. "Email (work)")
            for fld, syns in SYNONYMS.items():
                if fld not in taken and any(norm.startswith(s + " ") for s in syns):
                    match = fld
                    break
        if match:
            mapping[original] = match
            taken.add(match)
        elif original.strip():
            unmapped.append(original)
    return mapping, unmapped


@dataclass
class IngestResult:
    leads: list[RawLead]
    unmapped_columns: list[str]
    input_hash: str
    warnings: list[str] = field(default_factory=list)


def _read_rows(path: Path, data: bytes) -> list[dict[str, str]]:
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        if ws is None:
            return []
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = ["" if h is None else str(h) for h in rows[0]]
        return [
            {headers[i]: ("" if v is None else str(v)) for i, v in enumerate(r) if i < len(headers)}
            for r in rows[1:]
        ]
    text = data.decode("utf-8-sig", errors="replace")
    sample = text[:4096]
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return list(csv.DictReader(io.StringIO(text), dialect=dialect))


def ingest(path: Path | str) -> IngestResult:
    path = Path(path)
    data = path.read_bytes()
    input_hash = hashlib.sha256(data).hexdigest()
    rows = _read_rows(path, data)
    if not rows:
        return IngestResult([], [], input_hash, ["file has no data rows"])

    mapping, unmapped = map_headers(list(rows[0].keys()))
    warnings = []
    if "email" not in mapping.values():
        warnings.append("no email column found; every row will be rejected")

    leads: list[RawLead] = []
    for i, row in enumerate(rows, start=2):  # Row 1 is the header
        values: dict[str, str] = {}
        for original, fld in mapping.items():
            values[fld] = (row.get(original) or "").strip()
        if not any(values.values()):
            continue  # Skip blank rows
        if not values.get("full_name"):
            first, last = values.pop("first_name", ""), values.pop("last_name", "")
            values["full_name"] = f"{first} {last}".strip()
        leads.append(RawLead(row_number=i, **values))
    return IngestResult(leads, unmapped, input_hash, warnings)
